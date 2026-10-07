"""
Генерация PDF-отчёта по обращениям — для кнопки "Скачать отчёт" на
странице "Аналитика". Использует reportlab (чистый Python, без
внешних зависимостей вроде LibreOffice — работает и в Docker-
контейнере на Render так же, как и локально).

Кириллица требует отдельного TTF-шрифта — стандартные встроенные
шрифты reportlab (Helvetica и т.п.) кириллицу не поддерживают.
Шрифт (DejaVu Sans, свободная лицензия) лежит рядом, в app/assets/.
"""
import io
import os
from collections import Counter
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle
)
from sqlalchemy.orm import Session

from app.models.ticket import Ticket

_ASSETS_DIR = os.path.join(
    os.path.dirname(os.path.dirname(__file__)),
    "assets"
)

_NAVY = colors.HexColor("#1B3768")
_LIGHT_ROW = colors.HexColor("#F1EFE8")

_fonts_registered = False


def _ensure_fonts_registered() -> None:

    global _fonts_registered

    if _fonts_registered:
        return

    pdfmetrics.registerFont(
        TTFont("DejaVuSans", os.path.join(_ASSETS_DIR, "DejaVuSans.ttf"))
    )
    pdfmetrics.registerFont(
        TTFont(
            "DejaVuSans-Bold",
            os.path.join(_ASSETS_DIR, "DejaVuSans-Bold.ttf")
        )
    )

    _fonts_registered = True


def _styled_table(data: list[list[str]], col_widths=None) -> Table:

    table = Table(data, colWidths=col_widths, repeatRows=1)

    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "DejaVuSans"),
        ("FONTNAME", (0, 0), (-1, 0), "DejaVuSans-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("BACKGROUND", (0, 0), (-1, 0), _NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, _LIGHT_ROW]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D3D1C7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))

    return table


def _fetch_tickets_for_report(
    db: Session,
    date_from: date | None,
    date_to: date | None
) -> list[Ticket]:

    query = db.query(Ticket)

    if date_from:
        query = query.filter(Ticket.created_date >= date_from)

    if date_to:
        query = query.filter(Ticket.created_date <= date_to)

    return query.all()


def _compute_report_counts(tickets: list[Ticket]) -> dict:

    total = len(tickets)
    sent_count = sum(1 for t in tickets if t.email_sent)

    return {
        "total": total,
        "new_count": sum(1 for t in tickets if t.status == "Новое"),
        "in_progress": sum(1 for t in tickets if t.status == "В работе"),
        "closed": sum(1 for t in tickets if t.status == "Закрыто"),
        "sent_count": sent_count,
        "not_sent_count": total - sent_count,
        "category_counts": Counter(
            t.ai_category or "Не определено ИИ" for t in tickets
        ),
        "district_counts": Counter(
            (t.district or "Не указан").strip() or "Не указан"
            for t in tickets
        ),
        "operator_counts": Counter(
            t.operator or t.ai_operator or "Не определён" for t in tickets
        ),
    }


def generate_tickets_report_pdf(
    db: Session,
    date_from: date | None = None,
    date_to: date | None = None
) -> bytes:
    """
    Формирует PDF-отчёт по обращениям за указанный период (или за
    всё время, если период не задан). Возвращает готовые байты PDF.
    """

    _ensure_fonts_registered()

    tickets = _fetch_tickets_for_report(db, date_from, date_to)
    counts = _compute_report_counts(tickets)

    total = counts["total"]
    new_count = counts["new_count"]
    in_progress = counts["in_progress"]
    closed = counts["closed"]
    sent_count = counts["sent_count"]
    not_sent_count = counts["not_sent_count"]
    category_counts = counts["category_counts"]
    district_counts = counts["district_counts"]
    operator_counts = counts["operator_counts"]

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleRu",
        parent=styles["Title"],
        fontName="DejaVuSans-Bold",
        fontSize=19,
        textColor=_NAVY,
        spaceAfter=2
    )
    subtitle_style = ParagraphStyle(
        "SubtitleRu",
        parent=styles["Normal"],
        fontName="DejaVuSans",
        fontSize=10.5,
        textColor=colors.HexColor("#5F5E5A")
    )
    heading_style = ParagraphStyle(
        "HeadingRu",
        parent=styles["Heading2"],
        fontName="DejaVuSans-Bold",
        fontSize=13,
        textColor=_NAVY,
        spaceBefore=14,
        spaceAfter=6
    )

    if date_from and date_to:
        period_text = (
            f"Период: {date_from.strftime('%d.%m.%Y')} — "
            f"{date_to.strftime('%d.%m.%Y')}"
        )
    elif date_from:
        period_text = f"Период: с {date_from.strftime('%d.%m.%Y')}"
    elif date_to:
        period_text = f"Период: по {date_to.strftime('%d.%m.%Y')}"
    else:
        period_text = "Период: за всё время"

    elements = [
        Paragraph("Smart Aqmola — отчёт по обращениям", title_style),
        Paragraph(period_text, subtitle_style),
        Paragraph(
            f"Сформирован: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
            subtitle_style
        ),
        Spacer(1, 8 * mm),

        Paragraph("Сводка", heading_style),
        _styled_table(
            [
                ["Показатель", "Количество"],
                ["Всего обращений", str(total)],
                ["Новых", str(new_count)],
                ["В работе", str(in_progress)],
                ["Закрыто", str(closed)],
                ["Отправлено оператору", str(sent_count)],
                ["Не отправлено", str(not_sent_count)],
            ],
            col_widths=[110 * mm, 50 * mm]
        ),
    ]

    def _breakdown_section(title: str, counter: Counter) -> None:

        elements.append(Paragraph(title, heading_style))

        if not counter:
            elements.append(Paragraph("Нет данных за период.", subtitle_style))
            return

        rows = [["Значение", "Количество"]] + [
            [name, str(count)]
            for name, count in counter.most_common()
        ]

        elements.append(_styled_table(rows, col_widths=[110 * mm, 50 * mm]))

    _breakdown_section("По категориям (ИИ-анализ)", category_counts)
    _breakdown_section("По районам / городам", district_counts)
    _breakdown_section("По операторам связи", operator_counts)

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        title="Smart Aqmola — отчёт по обращениям"
    )

    doc.build(elements)

    buffer.seek(0)
    return buffer.getvalue()


_HEADER_FILL = PatternFill(
    start_color="1B3768",
    end_color="1B3768",
    fill_type="solid"
)
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def generate_tickets_report_xlsx(
    db: Session,
    date_from: date | None = None,
    date_to: date | None = None
) -> bytes:
    """
    То же самое, что и generate_tickets_report_pdf, только в формате
    Excel (.xlsx) — один лист "Сводка", всё подряд: сначала общие
    цифры, затем разбивки по категориям/районам/операторам —
    каждая как отдельная подгруппа строк на том же листе.
    """

    tickets = _fetch_tickets_for_report(db, date_from, date_to)
    counts = _compute_report_counts(tickets)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сводка"

    if date_from and date_to:
        period_text = (
            f"{date_from.strftime('%d.%m.%Y')} — "
            f"{date_to.strftime('%d.%m.%Y')}"
        )
    elif date_from:
        period_text = f"с {date_from.strftime('%d.%m.%Y')}"
    elif date_to:
        period_text = f"по {date_to.strftime('%d.%m.%Y')}"
    else:
        period_text = "за всё время"

    sheet.append(["Показатель", "Количество"])

    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")

    def _append_section_title(title: str) -> None:

        sheet.append([])
        sheet.append([title])

        title_row = sheet[sheet.max_row]
        title_row[0].font = Font(bold=True, color="1B3768")

    sheet.append(["Период", period_text])
    sheet.append(["Сформирован", datetime.now().strftime("%d.%m.%Y %H:%M")])
    sheet.append(["Всего обращений", counts["total"]])
    sheet.append(["Новых", counts["new_count"]])
    sheet.append(["В работе", counts["in_progress"]])
    sheet.append(["Закрыто", counts["closed"]])
    sheet.append(["Отправлено оператору", counts["sent_count"]])
    sheet.append(["Не отправлено", counts["not_sent_count"]])

    _append_section_title("По категориям (ИИ-анализ)")
    for name, count in counts["category_counts"].most_common():
        sheet.append([name, count])

    _append_section_title("По районам / городам")
    for name, count in counts["district_counts"].most_common():
        sheet.append([name, count])

    _append_section_title("По операторам связи")
    for name, count in counts["operator_counts"].most_common():
        sheet.append([name, count])

    sheet.column_dimensions[get_column_letter(1)].width = 42
    sheet.column_dimensions[get_column_letter(2)].width = 16

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()
